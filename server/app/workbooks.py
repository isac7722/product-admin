"""Fixed-template Excel previews, reviewed imports, and faithful XLSX export."""

import copy
import io
import re
import unicodedata
import warnings
import xml.etree.ElementTree as ET
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Annotated
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from fastapi import File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from app.models import Cost, Product, Rate, Shooting, Specification
from app.services import save_cost, save_product, save_rate, save_shooting
from app.storage import ROOT, all_docs, get_doc, save_doc, storage_dir, transaction


class Analyze(BaseModel):
    """Explicit current sheets and identity defaults selected by the user."""

    sheets: list[str] = Field(min_length=1)
    design: str = ""
    product_type: str = ""
    prefix: str = ""
    color: str = ""
    group_id: str = ""


class Choice(BaseModel):
    """A resolved source record; values may be corrected before applying."""

    key: str
    target_id: str | None = None
    values: dict
    fields: list[str]


class Apply(BaseModel):
    """Preview version prevents applying a stale diff twice."""

    version: int
    choices: list[Choice] = Field(min_length=1)


def text(value) -> str:
    """Convert only actual cell contents, preserving empty cells."""
    return "" if value is None else str(value).strip()


def number(value) -> str | None:
    """Never substitute zero for missing formulas or nonnumeric text."""
    if value is None or isinstance(value, bool):
        return None
    try:
        n = Decimal(str(value))
        return format(n, "f") if n.is_finite() and n >= 0 else None
    except InvalidOperation:
        return None


def normalized(value: str) -> str:
    """Normalization is for suggestions only, not automatic merging."""
    return re.sub(r"\s+", "", unicodedata.normalize("NFC", value)).casefold()


def read_books(path: Path):
    """Validate the archive before parsing potentially large sheets."""
    try:
        with ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 300 * 1024 * 1024:
                raise HTTPException(422, "압축 해제 크기가 너무 큰 엑셀입니다.")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            formulas = load_workbook(path, data_only=False, read_only=False)
            values = load_workbook(path, data_only=True, read_only=False)
        return formulas, values
    except (BadZipFile, KeyError, ValueError, OSError) as exc:
        raise HTTPException(422, "읽을 수 없는 XLSX 파일입니다.") from exc


def sheet_kind(sheet) -> str:
    """Recognize supported layouts by header content, not filename."""
    if text(sheet["A1"].value) == "제품 원가계산서" and text(sheet["B7"].value) == "품목":
        return (
            "costs" if "VAT포함" in text(sheet["C39"].value) and "원단가" in text(sheet["C17"].value) else "reference"
        )
    if "촬영구성" in text(sheet["A1"].value) and text(sheet["C8"].value) == "이불앞":
        return "shooting"
    if text(sheet["A1"].value) == "상품기술서" and "상품명" in text(sheet["A2"].value):
        return "spec"
    if sheet.title in RATE_LAYOUTS:
        header = " ".join(text(sheet.cell(r, c).value) for r in range(1, 4) for c in range(1, 15))
        if any(word in header for word in ("단가", "금액", "임가공")):
            return "rates"
    return "reference"


# first data row, name columns, amount column, category, unit, VAT included
RATE_LAYOUTS = {
    "원단 단가": (2, [1], 14, "원단", "yd", False),
    "커튼원단 단가": (2, [1], 14, "원단", "yd", False),
    "침구 공임 단가": (3, [1, 3, 5], 6, "제작공임", "개", False),
    "솜 공임 단가": (2, [1, 3, 5], 6, "제작공임", "개", False),
    "솜 단가": (5, [1, 2], 8, "솜", "yd", False),
    "부자재단가": (3, [2], 7, "부자재", "개", False),
    "택배비.포장비": (2, [1, 2], 3, "기타", "개", True),
    "소품 공임단가": (2, [1, 3, 5], 6, "제작공임", "개", False),
    "대량 소품 공임 단가": (2, [1, 3, 5], 6, "제작공임", "개", False),
    "중국 생산": (2, [1, 3, 5], 6, "기타", "개", False),
    "커튼 공임 단가": (2, [1, 3, 4], 5, "제작공임", "개", False),
}


def parse_rates(formulas, values) -> list[dict]:
    """Retain current rates and source formulas, not historical columns."""
    start, columns, amount_col, category, unit, vat = RATE_LAYOUTS[formulas.title]
    records = []
    for row in range(start, min(formulas.max_row + 1, 5000)):
        name = " / ".join(
            text(values.cell(row, col).value) for col in columns if values.cell(row, col).value is not None
        )
        original = formulas.cell(row, amount_col).value
        amount = number(values.cell(row, amount_col).value)
        if not name or "▼" in name or (original is None and amount is None):
            continue
        cat = text(values.cell(row, 1).value) if formulas.title == "택배비.포장비" else category
        if cat not in {"택배비", "포장비"} and formulas.title == "택배비.포장비":
            cat = "기타"
        is_formula = formulas.cell(row, amount_col).data_type == "f" or hasattr(original, "text")
        notes = []
        if amount is None:
            notes.append("단가 확인 필요: 수식 결과 또는 숫자가 없습니다.")
        if is_formula:
            notes.append("엑셀 저장 계산값입니다. 반영 전 단가를 확인하세요.")
        records.append(
            {
                "kind": "rates",
                "label": name,
                "values": {
                    "name": name,
                    "category": cat,
                    "unit": unit,
                    "amount": amount,
                    "vat_included": vat,
                    "conditions": " / ".join(
                        text(values.cell(row, c).value)
                        for c in ([4, 5, 10, 17] if category == "원단" else [4, 5])
                        if values.cell(row, c).value is not None
                    ),
                    "supplier": text(values.cell(row, 8 if category == "원단" else 2).value),
                    "notes": text(values.cell(row, 15 if category == "원단" else 8).value),
                },
                "cell": formulas.cell(row, amount_col).coordinate,
                "warnings": notes,
                "formula": getattr(original, "text", original) if is_formula else None,
            }
        )
    return records


def parse_spec(sheet, context: Analyze) -> list[dict]:
    """Extract the shared fields and explicit physical color/size candidates."""
    fields = {
        "name_ko": "B2",
        "name_en": "B3",
        "front_material": "C21",
        "back_material": "C22",
        "filling": "C23",
        "country": "C29",
        "notes": "B4",
    }
    spec = {name: text(sheet[cell].value) for name, cell in fields.items()}
    spec["keywords"] = "\n".join(text(sheet.cell(r, 2).value) for r in range(5, 9) if sheet.cell(r, 2).value)
    spec["description"] = "\n\n".join(
        f"{text(sheet.cell(r, 1).value)}\n{text(sheet.cell(r, 2).value)}"
        for r in range(31, min(sheet.max_row + 1, 100))
        if sheet.cell(r, 2).value
    )
    colors = []
    for col in (2, 4, 6):
        for row in range(11, 17):
            raw = text(sheet.cell(row, col).value)
            if raw:
                name = re.split(r"\s*(?:앞:|뒤:|#)", raw)[0].strip()
                if name and name not in colors:
                    colors.append(name)
    spec["color_details"] = "\n".join(
        text(sheet.cell(r, c).value) for r in range(11, 18) for c in range(2, 11) if sheet.cell(r, c).value
    )
    records = [{"kind": "spec", "label": spec["name_ko"], "values": spec, "cell": "A1:J50", "warnings": []}]
    sizes = []
    for row in range(25, 28):
        size = text(sheet.cell(row, 2).value)
        dimension = text(sheet.cell(row, 3).value)
        if size:
            if size == "K" and any(s[0] == "Q/K" for s in sizes):
                continue
            if size == "Q" and text(sheet["B27"].value) == "K" and text(sheet["C27"].value).startswith(dimension):
                size = "Q/K"
            sizes.append((size, dimension))
    for color in colors:
        for size, dimensions in sizes:
            records.append(
                {
                    "kind": "products",
                    "label": f"{color} / {size}",
                    "cell": "B11:J27",
                    "warnings": [],
                    "values": {
                        "design": context.design,
                        "product_type": context.product_type,
                        "prefix": context.prefix,
                        "color": color,
                        "size": size,
                        "dimensions": dimensions,
                        "notes": "",
                    },
                }
            )
    return records


def parse_shooting(sheet, context: Analyze) -> list[dict]:
    """Keep textual references until a user connects a physical product."""
    records = []
    for row in range(9, min(sheet.max_row + 1, 200)):
        if text(sheet.cell(row, 1).value) == "단종" or "제품 표현" in text(sheet.cell(row, 1).value):
            break
        color = text(sheet.cell(row, 2).value)
        if not color:
            continue
        records.append(
            {
                "kind": "shooting",
                "label": color,
                "cell": f"B{row}:I{row}",
                "warnings": [],
                "values": {
                    "color": color,
                    "front_color": text(sheet.cell(row, 3).value),
                    "back_color": text(sheet.cell(row, 4).value),
                    "items": [
                        {
                            "role": "기본 베개커버",
                            "product_id": None,
                            "notes": f"앞: {text(sheet.cell(row, 5).value)} / 뒤: {text(sheet.cell(row, 6).value)}",
                        },
                        {"role": "코디 베개커버", "product_id": None, "notes": text(sheet.cell(row, 7).value)},
                        {"role": "침대패드", "product_id": None, "notes": text(sheet.cell(row, 8).value)},
                    ],
                    "notes": f"매트리스커버: {text(sheet.cell(row, 9).value)}",
                },
            }
        )
    return records


def parse_costs(formulas, values, context: Analyze, rates: list[dict]) -> list[dict]:
    """Map production inputs to current rates without executing Excel formulas."""
    records = []
    for col, price_row in ((4, 17), (5, 10), (6, 13)):
        if not values.cell(7, col).value:
            continue
        issues = []
        lines = []

        def find_rate(name, category, quantity, loss="0", label="", issues=issues, lines=lines):
            candidates = [r for r in rates if r["category"] == category and normalized(r["name"]) == normalized(name)]
            if len(candidates) != 1 or quantity is None:
                issues.append(f"{label or name}: 단가 연결 또는 소요량 확인 필요")
                return
            lines.append(
                {"rate_id": candidates[0]["id"], "quantity": quantity, "loss_rate": loss, "label": label or name}
            )

        fabric_name = text(values.cell(11, col).value)
        quantity = number(values.cell(18, col).value)
        piping = number(values.cell(19, col).value)
        if quantity is not None:
            quantity = str(Decimal(quantity) + Decimal(piping or "0"))
        find_rate(fabric_name, "원단", quantity, "0.07", "원단")
        cotton = text(values.cell(22, col).value)
        if cotton:
            find_rate(
                f"{cotton} / {text(values.cell(23, col).value)}", "솜", number(values.cell(25, col).value), label="솜"
            )
        labor_name = " / ".join(text(values.cell(r, col).value) for r in (7, 8, 15))
        find_rate(labor_name, "제작공임", "1", label="제작공임")
        for row in range(28, 37):
            if formulas.cell(row, col).value is not None:
                find_rate(text(values.cell(row, 2).value), "부자재", "1")
        charges = {}
        for key, category, label_col in (("shipping_rate_id", "택배비", 16), ("packaging_rate_id", "포장비", 18)):
            name = f"{category} / {text(values.cell(price_row, label_col).value)}"
            candidates = [r for r in rates if r["category"] == category and normalized(r["name"]) == normalized(name)]
            charges[key] = candidates[0]["id"] if len(candidates) == 1 else None
            if not charges[key]:
                issues.append(f"{name}: 연결 확인 필요")
        records.append(
            {
                "kind": "costs",
                "label": f"{text(values.cell(7, col).value)} / {text(values.cell(15, col).value)}",
                "cell": f"{formulas.cell(7, col).coordinate}:{formulas.cell(39, col).coordinate}",
                "warnings": issues,
                "values": {
                    "lines": lines,
                    "unresolved_inputs": issues,
                    "sale_price": number(values.cell(price_row, 13).value),
                    **charges,
                },
                "identity": {"size": text(values.cell(15, col).value), "color": context.color},
            }
        )
    return records


def infer_candidates(db, record: dict, context: Analyze) -> list[dict]:
    """Offer matches for human confirmation, never merge on fuzzy names."""
    kind = record["kind"]
    if kind == "rates":
        return [r for r in all_docs(db, "rates") if normalized(r["name"]) == normalized(record["values"]["name"])]
    if kind in {"products", "costs"}:
        v = record.get("identity", record["values"])
        return [
            p
            for p in all_docs(db, "products")
            if (not context.design or p["design"] == context.design)
            and normalized(p["color"]) == normalized(v.get("color", ""))
            and (p["size"] == v.get("size") or p["size"] == "Q/K" and v.get("size") == "Q")
        ]
    return []


def register_routes(app, receive_file) -> None:
    """Attach source import and export endpoints to the local API."""

    @app.get("/api/imports")
    def imports():
        with transaction() as db:
            return [{k: v for k, v in x.items() if k not in {"records", "snapshots"}} for x in all_docs(db, "imports")][
                ::-1
            ]

    @app.post("/api/imports", status_code=201)
    async def upload(file: Annotated[UploadFile, File()], mode: Annotated[str, Form()] = "operation"):
        body = await receive_file(file, "excel")
        try:
            formulas, values = read_books(storage_dir() / "uploads" / body["path"])
            sheets = [{"name": s.title, "kind": sheet_kind(s)} for s in formulas]
            formulas.close()
            values.close()
            with transaction() as db:
                source = save_doc(db, "files", body)
                return save_doc(
                    db,
                    "imports",
                    {
                        "file_id": source["id"],
                        "name": body["name"],
                        "sheets": sheets,
                        "mode": mode,
                        "status": "시트 선택",
                        "records": [],
                        "snapshots": {},
                    },
                )
        except Exception:
            (storage_dir() / "uploads" / body["path"]).unlink(missing_ok=True)
            raise

    @app.get("/api/imports/{identifier}")
    def get_import(identifier: str):
        with transaction() as db:
            return get_doc(db, "imports", identifier)

    @app.post("/api/imports/{identifier}/analyze")
    def analyze(identifier: str, payload: Analyze):
        with transaction() as db:
            job = get_doc(db, "imports", identifier)
            if job["status"] == "반영 완료":
                raise HTTPException(409, "이미 반영한 파일입니다. 새 가져오기를 시작하세요.")
            source = get_doc(db, "files", job["file_id"])
            formulas, values = read_books(storage_dir() / "uploads" / source["path"])
            records = []
            for name in payload.sheets:
                if name not in formulas.sheetnames:
                    raise HTTPException(422, "선택한 시트를 찾을 수 없습니다.")
                sheet = formulas[name]
                kind = sheet_kind(sheet)
                if kind == "rates":
                    parsed = parse_rates(sheet, values[name])
                elif kind == "spec":
                    if text(values[name]["A21"].value) != "원단 및 소재":
                        raise HTTPException(422, f"{name}: 기술서 배치가 다릅니다. 현재 리뉴얼 양식을 선택하세요.")
                    parsed = parse_spec(values[name], payload)
                elif kind == "shooting":
                    parsed = parse_shooting(values[name], payload)
                elif kind == "costs":
                    parsed = parse_costs(sheet, values[name], payload, all_docs(db, "rates"))
                else:
                    raise HTTPException(422, f"{name}: 참고 시트는 원본 보관만 지원합니다.")
                for record in parsed:
                    key = f"{name}:{record['kind']}:{record['label']}:{record['cell']}"
                    candidates = infer_candidates(db, record, payload)
                    record.update(
                        key=key,
                        source={"file_id": job["file_id"], "sheet": name, "cell": record["cell"]},
                        candidates=candidates,
                        target_id=None,
                    )
                    mapping_key = f"{payload.design}|{payload.product_type}|{key}"
                    mapping = next((m for m in all_docs(db, "mappings") if m["key"] == mapping_key), None)
                    record["mapping_key"] = mapping_key
                    if mapping:
                        record["target_id"] = mapping["target_id"]
                    records.append(record)
            formulas.close()
            values.close()
            snapshots = {
                d["id"]: d["version"]
                for kind in ("products", "rates", "groups", "costs", "shooting")
                for d in all_docs(db, kind)
            }
            return save_doc(
                db,
                "imports",
                {
                    **job,
                    "status": "비교·선택",
                    "records": records,
                    "context": payload.model_dump(),
                    "snapshots": snapshots,
                },
                identifier,
                job["version"],
            )

    @app.post("/api/imports/{identifier}/apply")
    def apply(identifier: str, payload: Apply):
        with transaction() as db:
            job = get_doc(db, "imports", identifier)
            if job["status"] == "반영 완료":
                return job
            if job["version"] != payload.version:
                raise HTTPException(409, "분석 결과가 바뀌었습니다. 다시 확인하세요.")
            for resource in ("products", "rates", "groups", "costs", "shooting"):
                for item in all_docs(db, resource):
                    if item["id"] in job["snapshots"] and item["version"] != job["snapshots"][item["id"]]:
                        raise HTTPException(409, "비교 이후 값이 변경되었습니다. 다시 분석하세요.")
            records = {r["key"]: r for r in job["records"]}
            if len({c.key for c in payload.choices}) != len(payload.choices):
                raise HTTPException(422, "같은 행을 중복 반영할 수 없습니다.")
            group_id = job["context"].get("group_id")
            # Products precede shared specs so an empty installation can be migrated.
            choices = sorted(
                payload.choices, key=lambda c: 0 if records.get(c.key, {}).get("kind") == "products" else 1
            )
            applied = []
            for choice in choices:
                if choice.key not in records:
                    raise HTTPException(422, "분석되지 않은 행입니다.")
                record = records[choice.key]
                kind = record["kind"]
                target = choice.target_id
                old = None
                if kind == "costs" and target:
                    get_doc(db, "products", target)
                    old = next((c for c in all_docs(db, "costs") if c["product_id"] == target), None)
                elif kind in {"products", "rates"} and target:
                    old = get_doc(db, kind, target)
                elif kind == "spec" and group_id:
                    old = get_doc(db, "groups", group_id)
                elif kind == "shooting" and group_id:
                    old = next(
                        (
                            s
                            for s in all_docs(db, "shooting")
                            if s["group_id"] == group_id and s["color"] == choice.values.get("color")
                        ),
                        None,
                    )
                if old and old["id"] in job["snapshots"] and old["version"] != job["snapshots"][old["id"]]:
                    raise HTTPException(409, "비교 이후 값이 변경되었습니다. 다시 분석하세요.")
                if old and old["id"] not in job["snapshots"] and kind in {"products", "rates", "costs"}:
                    raise HTTPException(409, "비교 이후 대상이 생성되었습니다. 다시 분석하세요.")
                model = {
                    "products": Product,
                    "rates": Rate,
                    "costs": Cost,
                    "spec": Specification,
                    "shooting": Shooting,
                }[kind]
                original = old.get("spec", {}) if kind == "spec" and old else old or {}
                data = {k: v for k, v in original.items() if k in model.model_fields}
                for field in choice.fields:
                    if field not in model.model_fields or field in {"version", "source"}:
                        raise HTTPException(422, "반영할 수 없는 필드입니다.")
                    if field in choice.values:
                        data[field] = choice.values[field]
                        if data[field] == "" and field in {
                            "amount",
                            "sale_price",
                            "shipping_rate_id",
                            "packaging_rate_id",
                        }:
                            data[field] = None
                data.update(version=old["version"] if old else None, source=record["source"])
                if kind == "products" and "size" in choice.fields and "size_id" not in choice.fields:
                    data["size_id"] = None
                if kind == "products" and "dimensions" in choice.fields and "dimensions_id" not in choice.fields:
                    data["dimensions_id"] = None
                validated = model(**data)
                if kind == "products":
                    saved = save_product(db, validated, target)
                    group_id = group_id or saved["group_id"]
                elif kind == "rates":
                    saved = save_rate(db, validated, target, "엑셀 선택 반영")
                elif kind == "costs":
                    if not target:
                        raise HTTPException(422, "원가를 연결할 제품을 선택하세요.")
                    if record["warnings"] and len(validated.lines) < len(record["values"]["lines"]):
                        raise HTTPException(422, "원가 항목 연결을 확인하세요.")
                    saved = save_cost(db, target, validated, "엑셀 선택 반영")
                elif kind == "spec":
                    if not group_id:
                        raise HTTPException(422, "기술서를 연결할 제품을 먼저 선택하거나 함께 등록하세요.")
                    group = get_doc(db, "groups", group_id)
                    saved = save_doc(
                        db,
                        "groups",
                        {**group, "spec": validated.model_dump(mode="json", exclude={"version"})},
                        group_id,
                        group["version"],
                    )
                else:
                    if not group_id:
                        raise HTTPException(422, "촬영구성을 연결할 기술서를 선택하세요.")
                    saved = save_shooting(db, group_id, validated)
                if kind in {"products", "rates", "costs"}:
                    mapping = next((m for m in all_docs(db, "mappings") if m["key"] == record["mapping_key"]), None)
                    save_doc(
                        db,
                        "mappings",
                        {"key": record["mapping_key"], "target_id": target or saved["id"]},
                        mapping["id"] if mapping else None,
                        mapping["version"] if mapping else None,
                    )
                applied.append({"key": choice.key, "id": saved["id"], "fields": choice.fields})
            return save_doc(
                db, "imports", {**job, "status": "반영 완료", "applied": applied}, identifier, job["version"]
            )

    @app.get("/api/groups/{identifier}/export")
    def export(identifier: str, color: str = ""):
        with transaction() as db:
            group = get_doc(db, "groups", identifier)
            products = [
                p
                for p in all_docs(db, "products")
                if p["group_id"] == identifier and (not color or p["color"] == color)
            ]
            if not products:
                raise HTTPException(422, "출력할 제품이 없습니다.")
            costs = {c["product_id"]: c for c in all_docs(db, "costs")}
            data = export_workbook(group, products, costs)
            return Response(
                data,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": "attachment; filename=product-cost.xlsx"},
            )


def export_workbook(group: dict, products: list[dict], costs: dict) -> bytes:
    """Patch template XML to preserve drawings, merges, styles and print setup."""
    template = next(
        p for p in (ROOT / "data").glob("*.xlsx") if unicodedata.normalize("NFC", p.name).startswith("원가계산서-")
    )
    template_book = load_workbook(template, data_only=True, read_only=True)
    template_values = template_book.worksheets[0]
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    office_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ET.register_namespace("", ns)
    with ZipFile(template) as source:
        content = {name: source.read(name) for name in source.namelist()}
    book = ET.fromstring(content["xl/workbook.xml"])
    book.attrib.pop("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", None)
    sheets = book.find(f"{{{ns}}}sheets")
    first = sheets[0]
    relationships = ET.fromstring(content["xl/_rels/workbook.xml.rels"])
    first_rel = next(r for r in relationships if r.attrib["Id"] == first.attrib[f"{{{office_rel}}}id"])
    target = first_rel.attrib["Target"]
    template_path = target.lstrip("/") if target.startswith("/") else "xl/" + target
    original_sheet = ET.fromstring(content[template_path])
    # Export only newly generated current sheets; original uploads remain separately downloadable.
    sheets.clear()
    for rel in list(relationships):
        if rel.attrib["Type"].endswith(("/worksheet", "/externalLink", "/calcChain")):
            relationships.remove(rel)
    external = book.find(f"{{{ns}}}externalReferences")
    if external is not None:
        book.remove(external)
    names = book.find(f"{{{ns}}}definedNames")
    original_names = [
        copy.deepcopy(n) for n in (names if names is not None else []) if n.attrib.get("localSheetId") == "0"
    ]
    if names is not None:
        book.remove(names)
    names = ET.SubElement(book, f"{{{ns}}}definedNames")
    types = ET.fromstring(content["[Content_Types].xml"])
    type_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    pages = []
    for color in dict.fromkeys(p["color"] for p in products):
        variants = [p for p in products if p["color"] == color]
        for offset in range(0, len(variants), 3):
            block = variants[offset : offset + 3]
            evidence = [(p, d) for p in block for d in costs.get(p["id"], {}).get("details", [])]
            pages.extend((color, block, evidence[start : start + 28]) for start in range(0, max(1, len(evidence)), 28))
    for index, (color, variants, evidence) in enumerate(pages):
        sheet = copy.deepcopy(original_sheet)
        sheet.attrib.pop("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", None)
        cells = {c.attrib["r"]: c for c in sheet.iter(f"{{{ns}}}c")}

        def put(address, value, numeric=False, cells=cells):
            cell = cells.get(address)
            if cell is None:
                return
            for child in list(cell):
                cell.remove(child)
            cell.attrib.pop("t", None)
            if value is None:
                return
            if numeric:
                ET.SubElement(cell, f"{{{ns}}}v").text = str(value)
            else:
                cell.set("t", "inlineStr")
                ET.SubElement(ET.SubElement(cell, f"{{{ns}}}is"), f"{{{ns}}}t").text = str(value)

        # Clear every computed result so unsupported values never appear current.
        for address, cell in cells.items():
            if cell.find(f"{{{ns}}}f") is not None:
                put(address, None)
        for row in range(7, 70):
            for col in ("D", "E", "F"):
                put(f"{col}{row}", None)
        for row in range(10, 32):
            for col in "IJKLMNOPQRSTUVW":
                if row not in (24, 25):
                    put(f"{col}{row}", None)
        for row in range(26, 32):
            put(f"I{row}", "미산출")
        for address in cells:
            if int(re.sub(r"[A-Z]+", "", address)) >= 42:
                put(address, None)
        for address in ("H18", "H29", "M6", "W23"):
            put(address, None)
        put("A2", "출력일")
        put("B2", datetime.now(ZoneInfo("Asia/Seoul")).strftime("%Y.%m.%d"))
        put("A5", "출력 범위")
        put("A6", "금액 기준")
        put("A41", "현재 원가 근거 · 과거 기록은 업로드 원본에서 확인")
        for evidence_row, (evidence_product, detail) in enumerate(evidence, start=42):
            put(f"A{evidence_row}", "원가 근거")
            put(f"B{evidence_row}", evidence_product["size"])
            unit_price = detail.get("unit_price") or "미입력"
            amount = detail.get("amount") or "계산 확인 필요"
            put(
                f"C{evidence_row}",
                f"{detail['name']} / 단가 {unit_price} × 소요량 {detail['quantity']}"
                f" / 로스금액 {detail.get('loss_amount', '미입력')} / VAT포함 {amount}",
            )
        put("A1", "제품 원가계산서 · 최신 데이터")
        put("B3", f"{group['design']} · {group['product_type']}")
        put("B4", color)
        put("B5", "세트·도매 가격 미산출 / 계산 근거 미등록은 빈 금액으로 표시")
        put("B6", "원가: VAT 포함 / 모든 금액은 현재 저장값")
        put("C19", "로스금액(원)")
        put("C20", "원단합계 VAT포함(원)")
        put("C26", "솜합계 VAT포함(원)")
        put("C27", "공임 VAT포함(원)")
        put("C37", "부자재 VAT포함(원)")
        put("C38", "기타 VAT포함(원)")
        put("E2", len(variants), True)
        for col, row, p in zip(("D", "E", "F"), (10, 13, 17), variants, strict=False):
            cost = costs.get(p["id"], {})
            put(f"{col}7", p["product_type"])
            put(f"{col}15", p["size"])
            put(f"{col}16", p.get("dimensions", ""))
            put(f"{col}39", cost.get("production_cost"), True)
            put(f"I{row}", p["code"])
            for output_col, key in {
                "J": "production_cost",
                "K": "normal_price",
                "M": "sale_price",
                "N": "fee",
                "O": "shipping",
                "Q": "packaging",
                "S": "vat",
                "T": "defect",
                "U": "margin",
                "V": "margin_rate",
            }.items():
                put(f"{output_col}{row}", cost.get(key), True)
            put(f"L{row}", "0.4", True)
            details = cost.get("details", [])

            # Existing category rows retain their meaning; all displayed totals include VAT.
            def subtotal(category, details=details):
                found = [d for d in details if d.get("category") == category]
                if not found or any(d.get("amount") is None for d in found):
                    return None
                return str(sum((Decimal(d["amount"]) for d in found), Decimal(0)))

            fabrics = [d for d in details if d.get("category") == "원단"]
            if len(fabrics) == 1:
                put(f"{col}11", fabrics[0]["name"])
                put(f"{col}17", fabrics[0].get("unit_price"), True)
                put(f"{col}18", fabrics[0].get("quantity"), True)
                put(f"{col}19", fabrics[0].get("loss_amount"), True)
            elif fabrics:
                put(f"{col}11", " / ".join(d["name"] for d in fabrics))
            put(f"{col}20", subtotal("원단"), True)
            cotton = [d for d in details if d.get("category") == "솜"]
            if len(cotton) == 1:
                put(f"{col}22", cotton[0]["name"])
                put(f"{col}24", cotton[0].get("unit_price"), True)
                put(f"{col}25", cotton[0].get("quantity"), True)
            put(f"{col}26", subtotal("솜"), True)
            put(f"{col}27", subtotal("제작공임"), True)
            accessories = [d for d in details if d.get("category") == "부자재"]
            for expense_row in range(28, 37):
                name = text(template_values.cell(expense_row, 2).value)
                matched = [d for d in accessories if normalized(d["name"]) == normalized(name)]
                if matched and all(d.get("amount") is not None for d in matched):
                    put(f"{col}{expense_row}", str(sum((Decimal(d["amount"]) for d in matched), Decimal(0))), True)
            put(f"{col}37", subtotal("부자재"), True)
            put(f"{col}38", subtotal("기타"), True)
            if not cost or cost.get("status") != "계산 완료":
                put(f"{col}8", "계산 확인 필요")
        title = f"{color[:20]}-{index + 1}"
        title = re.sub(r"[\\/*?:\[\]]", "_", title)
        relationship_id = f"rIdCurrent{index + 1}"
        path = f"xl/worksheets/current{index + 1}.xml"
        content[path] = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
        rel_path = str(Path(template_path).parent / "_rels" / (Path(template_path).name + ".rels"))
        if rel_path in content:
            content[f"xl/worksheets/_rels/current{index + 1}.xml.rels"] = content[rel_path]
        ET.SubElement(
            sheets, f"{{{ns}}}sheet", {"name": title, "sheetId": str(index + 1), f"{{{office_rel}}}id": relationship_id}
        )
        ET.SubElement(
            relationships,
            f"{{{rel_ns}}}Relationship",
            {"Id": relationship_id, "Type": f"{office_rel}/worksheet", "Target": f"worksheets/current{index + 1}.xml"},
        )
        ET.SubElement(
            types,
            f"{{{type_ns}}}Override",
            {
                "PartName": "/" + path,
                "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
            },
        )
        for original in original_names:
            item = copy.deepcopy(original)
            item.set("localSheetId", str(index))
            item.text = re.sub(r"^'[^']+'!", f"'{title}'!", item.text or "")
            names.append(item)
    book_views = book.find(f"{{{ns}}}bookViews")
    if book_views is not None:
        for view in book_views:
            view.set("activeTab", "0")
    content["xl/workbook.xml"] = ET.tostring(book, encoding="utf-8", xml_declaration=True)
    content["xl/_rels/workbook.xml.rels"] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
    content["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    template_book.close()
    result = io.BytesIO()
    with ZipFile(result, "w") as target:
        for name, data in content.items():
            target.writestr(name, data)
    return result.getvalue()
