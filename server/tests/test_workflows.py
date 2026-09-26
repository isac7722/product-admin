"""Acceptance tests against isolated databases and original reference workbooks."""

import hashlib
import io
import unicodedata
from decimal import Decimal
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from PIL import Image

from app.main import app
from app.storage import ROOT, transaction


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("PRODUCT_ADMIN_STORAGE", str(tmp_path))
    with TestClient(app) as client:
        yield client


def post(client, url, body):
    response = client.post(url, json=body)
    assert response.status_code in (200, 201), response.text
    return response.json()


def put(client, url, body):
    response = client.put(url, json=body)
    assert response.status_code == 200, response.text
    return response.json()


def product(client, **changes):
    size = changes.get("size", "SS")
    if not any(s["name"] == size for s in client.get("/api/sizes").json()):
        post(client, "/api/sizes", {"name": size})
    return post(
        client,
        "/api/products",
        {
            "design": "도즈",
            "product_type": "사계절 차렵이불",
            "prefix": "DOZE01",
            "color": "화이트",
            "size": "SS",
            **changes,
        },
    )


def rate(client, **changes):
    return post(
        client, "/api/rates", {"name": "60수 원단", "category": "원단", "unit": "yd", "amount": "1000", **changes}
    )


def cost(client, p, r, **changes):
    shipping = rate(client, name="택배비 소", category="택배비", amount="2000", vat_included=True)
    packaging = rate(client, name="포장비 소", category="포장비", amount="500", vat_included=True)
    return put(
        client,
        f"/api/products/{p['id']}/cost",
        {
            "lines": [{"rate_id": r["id"], "quantity": "2.5", "loss_rate": "0.07"}],
            "sale_price": "70000",
            "shipping_rate_id": shipping["id"],
            "packaging_rate_id": packaging["id"],
            **changes,
        },
    )


def editable(item, model):
    return {k: v for k, v in item.items() if k in model.model_fields}


def test_identity_and_prefix_propagation(client):
    from app.models import Product

    first = product(client)
    second = product(client, size="Q/K")
    assert first["group_id"] == second["group_id"]
    assert (
        client.post(
            "/api/products",
            json={
                "design": "도즈",
                "product_type": "사계절 차렵이불",
                "prefix": "DOZE01",
                "color": "화이트",
                "size": "SS",
            },
        ).status_code
        == 409
    )
    updated = put(
        client, f"/api/products/{first['id']}", {**editable(first, Product), "prefix": "NEW", "color": "크림"}
    )
    assert updated["code"] == "NEW-크림-SS"
    sibling = client.get(f"/api/products/{second['id']}").json()
    assert sibling["code"] == "NEW-크림-Q/K"
    assert client.put(f"/api/products/{first['id']}", json=editable(first, Product)).status_code == 409
    assert product(client, product_type="간절기 차렵이불", prefix="DOZE02")["group_id"] != first["group_id"]


def test_cost_decimal_history_and_atomicity(client):
    from app.models import Rate

    p = product(client)
    r = rate(client)
    result = cost(client, p, r)
    assert Decimal(result["production_cost"]) == Decimal("2942.5")
    assert Decimal(result["details"][0]["loss_amount"]) == Decimal("175")
    assert result["normal_price"] == "116700"
    expected = (Decimal(70000) - Decimal("2942.5") - 14000 - 2000 - 500) * Decimal("0.9") - 700
    assert Decimal(result["margin"]) == expected
    put(client, f"/api/rates/{r['id']}", {**editable(r, Rate), "amount": "2000"})
    changed = client.get(f"/api/products/{p['id']}/cost").json()
    assert Decimal(changed["production_cost"]) == 5885
    assert changed["sale_price"] == "70000" and changed["normal_price"] == "116700"
    history = client.get(f"/api/history/{p['id']}").json()
    assert len(history) == 1 and "sale_price" not in history[0]["changes"]
    current_rate = next(x for x in client.get("/api/rates").json() if x["id"] == r["id"])
    # An invalid category affects an existing production dependency; every write must roll back.
    assert (
        client.put(f"/api/rates/{r['id']}", json={**editable(current_rate, Rate), "category": "택배비"}).status_code
        == 422
    )
    assert next(x for x in client.get("/api/rates").json() if x["id"] == r["id"])["category"] == "원단"


def test_missing_zero_and_shipping_dependency(client):
    from app.models import Cost, Rate

    p = product(client)
    r = rate(client, amount=None)
    missing = cost(client, p, r)
    assert missing["production_cost"] is None and missing["margin"] is None
    put(client, f"/api/rates/{r['id']}", {**editable(r, Rate), "amount": "0"})
    c = client.get(f"/api/products/{p['id']}/cost").json()
    zero = put(client, f"/api/products/{p['id']}/cost", {**editable(c, Cost), "sale_price": "0"})
    assert zero["normal_price"] == "0" and zero["margin_rate"] is None
    assert Decimal(zero["margin"]) < 0
    shipping = next(x for x in client.get("/api/rates").json() if x["id"] == c["shipping_rate_id"])
    put(client, f"/api/rates/{shipping['id']}", {**editable(shipping, Rate), "amount": "3000"})
    assert client.get(f"/api/products/{p['id']}/cost").json()["shipping"] == "3000"


def test_shared_spec_shooting_sales(client):
    p = product(client)
    other = product(client, design="다른 디자인", product_type="베개커버", prefix="OTHER")
    g = client.get(f"/api/groups/{p['group_id']}").json()
    put(
        client,
        f"/api/groups/{g['id']}/spec",
        {"version": g["version"], "name_ko": "테스트 기술서", "front_material": "면 100%"},
    )
    put(
        client,
        f"/api/groups/{g['id']}/shooting",
        {"color": "화이트", "items": [{"role": "코디 베개커버", "product_id": other["id"], "notes": "뒤집어서 촬영"}]},
    )
    sale = post(
        client,
        "/api/sales",
        {
            "group_id": g["id"],
            "channel": "자사몰",
            "sale_type": "세트",
            "name": "화이트 퀸 세트",
            "items": [{"product_id": p["id"], "quantity": 1}, {"product_id": other["id"], "quantity": 2}],
        },
    )
    option = post(
        client,
        "/api/sales",
        {
            "group_id": g["id"],
            "channel": "자사몰",
            "sale_type": "추가옵션",
            "name": "추가 베개",
            "parent_id": sale["id"],
            "items": [{"product_id": other["id"], "quantity": 1}],
        },
    )
    assert option["parent_id"] == sale["id"]
    g = client.get(f"/api/groups/{g['id']}").json()
    assert g["spec"]["name_ko"] == "테스트 기술서" and len(g["sales"]) == 2 and len(g["shooting"]) == 1


def test_upload_replace_original_integrity(client):
    p = product(client)
    image = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(image, "JPEG")
    data = image.getvalue()
    fields = {"group_id": p["group_id"], "purpose": "color", "color": "화이트"}
    first = client.post("/api/files", data=fields, files={"file": ("한글.jpg", data, "image/jpeg")})
    assert first.status_code == 201, first.text
    f = first.json()
    assert client.get(f"/api/files/{f['id']}/download").content == data
    assert client.get(f"/api/files/{f['id']}/preview").status_code == 200
    assert client.post("/api/files", data=fields, files={"file": ("new.jpg", data)}).status_code == 409
    assert (
        client.post(
            "/api/files", data={**fields, "replace_id": f["id"]}, files={"file": ("bad.jpg", b"not-image")}
        ).status_code
        == 422
    )
    assert len(client.get("/api/files").json()) == 1
    replacement = client.post("/api/files", data={**fields, "replace_id": f["id"]}, files={"file": ("new.jpg", data)})
    assert replacement.status_code == 201
    assert len(client.get("/api/files").json()) == 1
    assert client.get(f"/api/files/{f['id']}/download").content == data
    psd = b"8BPS\x00\x01" + b"\0" * 30
    response = client.post(
        "/api/files", data={"group_id": p["group_id"], "purpose": "psd"}, files={"file": ("source.psd", psd)}
    )
    assert response.status_code == 201 and client.get(f"/api/files/{response.json()['id']}/download").content == psd


def reference(prefix):
    return next(p for p in (ROOT / "data").glob("*.xlsx") if unicodedata.normalize("NFC", p.name).startswith(prefix))


def upload(client, prefix):
    source = reference(prefix)
    response = client.post(
        "/api/imports", data={"mode": "migration"}, files={"file": (source.name, source.read_bytes())}
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_real_rates_selective_import_mapping_and_stale(client):
    from app.models import Rate

    job = upload(client, "자재")
    data = client.get(f"/api/files/{job['file_id']}/download").content
    assert hashlib.sha256(data).hexdigest() == hashlib.sha256(reference("자재").read_bytes()).hexdigest()
    analyzed = post(client, f"/api/imports/{job['id']}/analyze", {"sheets": ["택배비.포장비"]})
    r = analyzed["records"][0]
    body = {
        "version": analyzed["version"],
        "choices": [{"key": r["key"], "values": r["values"], "fields": list(r["values"])}],
    }
    assert client.get("/api/rates").json() == []
    result = post(client, f"/api/imports/{job['id']}/apply", body)
    assert result["status"] == "반영 완료"
    post(client, f"/api/imports/{job['id']}/apply", body)
    assert len(client.get("/api/rates").json()) == 1
    imported = client.get("/api/rates").json()[0]
    assert imported["amount"] == "2100" and imported["source"]["cell"] == "C2"
    # A new upload reuses the confirmed alias, but comparison alone never writes.
    second = upload(client, "자재")
    analyzed = post(client, f"/api/imports/{second['id']}/analyze", {"sheets": ["택배비.포장비"]})
    r = analyzed["records"][0]
    assert r["target_id"] == imported["id"]
    put(client, f"/api/rates/{imported['id']}", {**editable(imported, Rate), "amount": "2300"})
    response = client.post(
        f"/api/imports/{second['id']}/apply",
        json={
            "version": analyzed["version"],
            "choices": [{"key": r["key"], "target_id": imported["id"], "values": r["values"], "fields": ["amount"]}],
        },
    )
    assert response.status_code == 409
    assert client.get("/api/rates").json()[0]["amount"] == "2300"


def test_real_spec_and_cost_preview(client):
    job = upload(client, "상품")
    analyzed = post(
        client,
        f"/api/imports/{job['id']}/analyze",
        {
            "sheets": ["차렵이불-리뉴얼", "촬영구성"],
            "design": "도즈",
            "product_type": "사계절 차렵이불",
            "prefix": "DOZE01",
        },
    )
    records = analyzed["records"]
    products = [r for r in records if r["kind"] == "products"]
    assert len(products) == 30 and {p["values"]["size"] for p in products} == {"SS", "Q/K"}
    selected = [r for r in records if r["kind"] == "spec"] + products[:2]
    for name in {r["values"]["size"] for r in products[:2]}:
        post(client, "/api/sizes", {"name": name})
    for name in {r["values"]["dimensions"] for r in products[:2] if r["values"]["dimensions"]}:
        post(client, "/api/dimensions", {"name": name})
    result = post(
        client,
        f"/api/imports/{job['id']}/apply",
        {
            "version": analyzed["version"],
            "choices": [{"key": r["key"], "values": r["values"], "fields": list(r["values"])} for r in selected],
        },
    )
    assert len(result["applied"]) == 3
    assert client.get("/api/products").json()["total"] == 2
    job = upload(client, "원가")
    analyzed = post(
        client,
        f"/api/imports/{job['id']}/analyze",
        {"sheets": ["260326 솜단가인상"], "design": "도즈", "color": "화이트"},
    )
    assert len(analyzed["records"]) == 3
    assert analyzed["records"][0]["values"]["unresolved_inputs"]


def test_export_layout_and_supported_values(client):
    p = product(client)
    r = rate(client)
    c = cost(client, p, r)
    result = client.get(f"/api/groups/{p['group_id']}/export")
    assert result.status_code == 200
    actual = load_workbook(io.BytesIO(result.content), data_only=True)
    original = load_workbook(reference("원가"), data_only=True)
    sheet = actual.worksheets[0]
    assert list(sheet.merged_cells.ranges) == list(original.worksheets[0].merged_cells.ranges)
    assert sheet.page_setup.orientation == original.worksheets[0].page_setup.orientation
    assert sheet.print_area.split("!")[-1] == original.worksheets[0].print_area.split("!")[-1]
    assert Decimal(str(sheet["J10"].value)) == Decimal(c["production_cost"])
    assert sheet["K10"].value == 116700 and sheet["M10"].value == 70000
    assert sheet["M11"].value is None and sheet["O26"].value is None
    assert sheet["I26"].value == "미산출"
    with ZipFile(io.BytesIO(result.content)) as z, ZipFile(reference("원가")) as original_zip:
        assert z.read("xl/styles.xml") == original_zip.read("xl/styles.xml")


def test_transaction_rollback(client):
    with pytest.raises(RuntimeError), transaction() as db:
        db.execute("INSERT INTO history VALUES('rollback','x','now','test','{}')")
        raise RuntimeError("injected failure")
    assert client.get("/api/history/x").json() == []


def test_rename_keeps_shared_color_assets_and_shooting(client):
    from app.models import Product

    p = product(client)
    sibling = product(client, size="Q/K")
    put(client, f"/api/groups/{p['group_id']}/shooting", {"color": "화이트", "notes": "촬영 메모"})
    img = io.BytesIO()
    Image.new("RGB", (10, 10)).save(img, "JPEG")
    upload_result = client.post(
        "/api/files",
        data={"group_id": p["group_id"], "purpose": "color", "color": "화이트"},
        files={"file": ("color.jpg", img.getvalue())},
    )
    assert upload_result.status_code == 201
    put(client, f"/api/products/{p['id']}", {**editable(p, Product), "color": "크림", "design": "도즈 리뉴얼"})
    group = client.get(f"/api/groups/{p['group_id']}").json()
    assert group["files"][0]["color"] == "크림"
    assert group["shooting"][0]["color"] == "크림"
    assert next(v for v in group["products"] if v["id"] == sibling["id"])["design"] == "도즈 리뉴얼"


def test_select_one_field_preserves_unselected_and_conflicting_batch_rolls_back(client):
    existing = rate(client, name="택배비 / 극소", category="택배비", amount="9999", notes="사이트에서 관리한 비고")
    job = upload(client, "자재")
    preview = post(client, f"/api/imports/{job['id']}/analyze", {"sheets": ["택배비.포장비"]})
    first = preview["records"][0]
    post(
        client,
        f"/api/imports/{job['id']}/apply",
        {
            "version": preview["version"],
            "choices": [
                {"key": first["key"], "target_id": existing["id"], "values": first["values"], "fields": ["amount"]}
            ],
        },
    )
    current = next(r for r in client.get("/api/rates").json() if r["id"] == existing["id"])
    assert current["amount"] == "2100" and current["notes"] == "사이트에서 관리한 비고"
    job = upload(client, "자재")
    preview = post(client, f"/api/imports/{job['id']}/analyze", {"sheets": ["택배비.포장비"]})
    first, second = preview["records"][:2]
    response = client.post(
        f"/api/imports/{job['id']}/apply",
        json={
            "version": preview["version"],
            "choices": [
                {
                    "key": first["key"],
                    "target_id": existing["id"],
                    "values": {**first["values"], "amount": "3000"},
                    "fields": ["amount"],
                },
                {
                    "key": second["key"],
                    "values": {**second["values"], "amount": "-1"},
                    "fields": list(second["values"]),
                },
            ],
        },
    )
    assert response.status_code == 422
    assert client.get("/api/rates").json()[0]["amount"] == "2100"
    assert client.get(f"/api/imports/{job['id']}").json()["status"] == "비교·선택"


def test_unresolved_cost_never_becomes_valid_and_export_splits_colors(client):
    p = product(client)
    product(client, color="크림")
    r = rate(client)
    result = cost(client, p, r, unresolved_inputs=["외부 침대패드 근거 누락"])
    assert result["production_cost"] is None and result["margin"] is None
    response = client.get(f"/api/groups/{p['group_id']}/export")
    book = load_workbook(io.BytesIO(response.content), data_only=True)
    assert len(book.sheetnames) == 2
    assert book.worksheets[0]["J10"].value is None
    assert book.worksheets[1]["B4"].value == "크림"


def test_storage_survives_application_restart(client):
    p = product(client)
    with TestClient(app) as restarted:
        assert restarted.get(f"/api/products/{p['id']}").json()["code"] == p["code"]
        assert restarted.get("/api/health").json()["status"] == "ok"
